"""Every result of the project in one table: one result per row, one value per cell (Excel + PDF + CSV).

v3 rows (multiple-choice accuracy, evidence hits, open-ended correctness) are computed from the saved result files
(runs/v3_baseline/rows.jsonl, runs/v3_baseline/open_judged.jsonl), with paired 95% bootstrap intervals over videos.
v1, v2, the v3 benchmark checks and the Head A / Head B method screens were measured by other scripts; their rows are
transcribed from the committed reports (docs/final_report_controller_study.md, reports/benchmark_v1*, docs/v2/*.md,
docs/v3/gate_*.md, docs/v3/research_log.md steps 29-45).

    python scripts/build_results_table.py      # docs/report/all_results.{xlsx,csv,html,pdf}
"""

from __future__ import annotations

import collections
import csv
import json
import subprocess
from pathlib import Path

import numpy as np

OUT = Path("docs/report")
RUN = Path("runs/v3_baseline")
B = 10_000                     # as declared in docs/v3/final_test_protocol.md


def family(ds: str) -> str:
    """Bootstrap cluster namespace: both Video-MMMU tracks use the same lecture videos, so they share one cluster."""
    return "videommmu" if ds.startswith("videommmu") else ds.replace("_test", "")

COLS = ["No.", "Stage", "Experiment", "System (plain name)", "Speech path (Path A)", "Picture search (Path B)",
        "Picker of the final pictures", "Pictures shown", "Text given to the answer model", "Answer model",
        "Trained parts used", "Benchmark", "Video type", "Question set", "Questions", "Videos", "What is measured",
        "Result (%)", "Correct / total", "vs starting framework: difference (points)",
        "vs starting framework: 95% interval", "vs starting framework: confirmed?",
        "vs best framework: difference (points)", "vs best framework: 95% interval", "vs best framework: confirmed?",
        "What this row means"]

STAGES = {1: "1. Version 1", 2: "2. Version 2", 3: "3. v3 benchmark check", 4: "4. v3 tuning (dev)",
          5: "5. Head A", 6: "6. Head B", 7: "7. Final locked test", 8: "8. Speech text", 9: "9. Larger answer model (4B)",
          10: "10. Full combination grid", 11: "11. Pooled summaries"}

DATA = {  # dataset key: (benchmark, video type, question set)
    "cgbench": ("CG-Bench", "long everyday videos (10-80 min)", "dev"),
    "cgbench_test": ("CG-Bench", "long everyday videos (10-80 min)", "test"),
    "videommmu": ("Video-MMMU Perception", "lectures, screen-reading questions", "dev"),
    "videommmu_test": ("Video-MMMU Perception", "lectures, screen-reading questions", "test"),
    "videommmu_comp": ("Video-MMMU Comprehension", "lectures, understanding questions", "dev"),
    "videommmu_comp_test": ("Video-MMMU Comprehension", "lectures, understanding questions", "test"),
    "ALL_TEST": ("CG-Bench + Video-MMMU (both tracks)", "all test questions", "test"),
    "LECTURE_TEST": ("Video-MMMU (both tracks)", "all lecture test questions", "test"),
}
POOL = {  # pool: (speech path, picture search, pictures in pool note)
    "v2": ("speech rule: MiniLM on lines inside the 4 BM25 windows", "scan every 5 s inside the BM25 windows"),
    "balanced": ("speech rule: MiniLM on lines inside the 4 BM25 windows", "scan inside the windows, spread evenly"),
    "video": ("speech rule: MiniLM on lines inside the 4 BM25 windows", "whole video, evenly (48 pictures)"),
    "dense": ("speech rule: MiniLM on lines inside the 4 BM25 windows", "dense whole-video search (picture every 2 s)"),
    "hybrid": ("speech rule: MiniLM on lines inside the 4 BM25 windows", "hybrid: 3 inside windows + 3 dense whole video"),
    "hybrid_a3": ("speech rule, only 3 pictures", "hybrid: 3 inside windows + 3 dense whole video"),
    "hybrid_b12": ("speech rule: MiniLM on lines inside the 4 BM25 windows", "hybrid, 12 pictures instead of 6"),
    "hybrid_azs": ("MiniLM over the whole transcript (untrained)", "hybrid: 3 inside windows + 3 dense whole video"),
    "hybrid_ahead": ("Head A (trained) over the whole transcript", "hybrid: 3 inside windows + 3 dense whole video"),
    "uniform4": ("none", "4 evenly spaced pictures"),
    "dense4": ("none", "dense whole-video search, top 4"),
}
PICK = {
    "clip": "top pictures by match to the question", "clipopt": "top pictures by match to question + options",
    "mmr": "MMR rule on the question", "mmropt": "MMR rule on question + options (best framework's picker)",
    "relevance": "top pictures by speech score", "A_only": "only speech-path pictures (best by question + options)",
    "B_only": "only picture-search pictures (best by question + options)", "none": "no pictures",
    "optset": "option-evidence set rule (untrained)", "ev": "Head B trained on human-marked seconds, no option term",
    "ev_set": "Head B (trained, chosen version)", "ev_set_noset": "Head B without avoiding look-alikes",
    "util": "Head B trained on answer-model feedback", "util_set": "Head B feedback version + option term",
}


def sig(lo, hi):
    return "yes (gain)" if lo > 0 else ("yes (loss)" if hi < 0 else "no")


def paired(a: dict, b: dict):
    qs = sorted(set(a) & set(b))
    if len(qs) < 10:
        return None
    byv = collections.defaultdict(list)
    for q in qs:
        byv[a[q]["_v"]].append(float(a[q]["correct"]) - float(b[q]["correct"]))
    vids = sorted(byv)
    rng = np.random.default_rng(20261019)
    boots = [np.mean(np.concatenate([byv[vids[i]] for i in rng.choice(len(vids), len(vids))])) for _ in range(B)]
    d = 100 * np.mean([float(a[q]["correct"]) - float(b[q]["correct"]) for q in qs])
    lo, hi = 100 * np.percentile(boots, 2.5), 100 * np.percentile(boots, 97.5)
    return round(d, 1), f"{lo:+.1f} to {hi:+.1f}", sig(lo, hi)


def describe(pool, selector, words):
    base = selector.split("|")[0]
    tags = selector.split("|")[1:]
    tp = next((t.split("=")[1] for t in tags if t.startswith("tp=")), None)
    model = "Qwen3-VL-4B" if "ans=4b" in tags else "Qwen3-VL-2B"
    if pool in ("verify", "verify6"):
        n = 4 if pool == "verify" else 6
        rule = "most confident answer" if base == "max" else "sum of option probabilities"
        return dict(path_a="speech rule and Head A (both pools)", path_b="hybrid", picker=f"confidence choice over {n} picture sets ({rule})",
                    text="none", model=model, trained="Head A, Head B" if n == 6 else "Head A", pics="4 (per set)",
                    name=f"Best framework + confidence choice ({n} sets, {'most confident answer' if base == 'max' else 'summed probabilities'})")
    if base.startswith("file:"):
        key = base.split("_hybrid_ahead_")[-1] if "_hybrid_ahead_" in base else base.split("_hybrid_")[-1]
    else:
        key = base
    path_a, path_b = POOL.get(pool, (pool, pool))
    picker = PICK.get(key, key)
    if words <= 0:
        text = "none"
    elif tp is None:
        text = f"BM25 excerpt, {words} words"
    else:
        text = {"frames+hybrid": f"speech around the chosen pictures + best lines (word + meaning search), {words} words",
                "hybrid": f"best lines by word + meaning search, {words} words",
                "frames": f"speech around the chosen pictures, {words} words",
                "full": f"whole transcript, capped at {words} words"}.get(tp, tp)
    trained = ", ".join(x for x in (("Head A" if pool == "hybrid_ahead" else ""), ("Head B" if key in ("ev", "ev_set", "ev_set_noset", "util", "util_set") else "")) if x) or "none"
    pics = {"none": "0"}.get(key, "4")
    if pool == "v2" and key == "clip" and words == 120 and tp is None:
        name = "Starting framework (v2 design)"
    elif pool == "hybrid" and key == "mmropt":
        name = "Best framework" + ("" if text == "none" else " + speech text" if tp == "frames+hybrid" else f" + {text}")
    elif pool == "hybrid_ahead" and key == "mmropt":
        name = "Best framework + Head A" + (" + speech text" if tp == "frames+hybrid" else "")
    elif pool == "hybrid" and key == "ev_set":
        name = "Best framework + Head B" + (" + speech text" if tp == "frames+hybrid" else "")
    elif pool == "hybrid_ahead" and key == "ev_set":
        name = "Best framework + Head A + Head B" + (" + speech text" if tp == "frames+hybrid" else "")
    elif pool == "uniform4":
        name = "Standard rule: 4 evenly spaced pictures"
    elif pool == "dense4":
        name = "Standard rule: dense MobileCLIP top-4"
    elif key == "none":
        name = "No pictures (question and options only)"
    else:
        name = f"Tuning variant: {path_b}; {picker}"
    return dict(path_a=path_a, path_b=path_b, picker=picker, text=text, model=model, trained=trained, pics=pics, name=name)


def stage_of(ds, pool, selector, words):
    tags = selector.split("|")[1:]
    four_b = "ans=4b" in tags
    tp = any(t.startswith("tp=") for t in tags)
    head = pool == "hybrid_ahead" or "ev_set" in selector or "_util" in selector or "_optset" in selector or "_ev" in selector
    test = ds.endswith("_test") or ds in ("ALL_TEST", "LECTURE_TEST")
    if four_b:
        return 10 if (head and (pool == "hybrid_ahead" or "ev_set" in selector)) else 9
    if pool in ("verify", "verify6"):
        return 7 if test else 6
    if test:
        if tp:
            return 10 if head else 8
        return 7
    if tp or (pool == "hybrid" and words in (120, 300) and selector == "mmropt" and ds == "videommmu_comp"):
        return 10 if head and tp else 8
    if selector.startswith("file:"):
        return 6
    if pool in ("hybrid_azs", "hybrid_ahead"):
        return 5
    return 4


def load_rows():
    last = {}
    for line in (RUN / "rows.jsonl").read_text(encoding="utf-8").splitlines():
        if line.strip():
            r = json.loads(line)
            if r["dataset"] in DATA:
                last[(r["dataset"], r["pool"], r["selector"], r["k"], r["words"], r["qa_id"])] = r
    groups = collections.defaultdict(dict)
    for (d, p, s, k, w, q), r in last.items():
        if k == 4:
            groups[(d, p, s, w)][(d, q)] = {**r, "_v": family(d) + ":" + r["video_id"]}
    return groups


def v3_rows(groups):
    out = []

    def ref_start(ds, sel):
        return groups.get((ds, "v2", "clip" + ("|ans=4b" if "ans=4b" in sel else ""), 120))

    def ref_best_key(ds, sel, words):
        """Best framework with the same answer model; with the 600-word speech text if this row uses it, else none."""
        tags = sel.split("|")[1:]
        m4 = "|ans=4b" if "ans=4b" in tags else ""
        if "tp=frames+hybrid" in tags and words == 600:
            return (ds, "hybrid", "mmropt|tp=frames+hybrid" + m4, 600)
        return (ds, "hybrid", "mmropt" + m4, 0)

    def ref_best(ds, sel, words):
        return groups.get(ref_best_key(ds, sel, words))

    def ref_start_key(ds, sel):
        return (ds, "v2", "clip" + ("|ans=4b" if "ans=4b" in sel else ""), 120)

    # pooled pseudo-datasets for test configurations present on all three test sets
    tests = ("cgbench_test", "videommmu_test", "videommmu_comp_test")
    for (d, p, s, w), g in list(groups.items()):
        if d != "cgbench_test":
            continue
        s_other = {t: s.replace("cgbench_test", t) for t in tests}
        if all((t, p, s_other[t], w) in groups for t in tests):
            allg, lec = {}, {}
            for t in tests:
                allg.update(groups[(t, p, s_other[t], w)])
                if t != "cgbench_test":
                    lec.update(groups[(t, p, s_other[t], w)])
            groups[("ALL_TEST", p, s.replace("cgbench_test", "ALL"), w)] = allg
            groups[("LECTURE_TEST", p, s.replace("cgbench_test", "LECTURE"), w)] = lec
    for (d, p, s, w), g in groups.items():
        desc = describe(p, s, w)
        bench, vtype, qset = DATA[d]
        n = len(g)
        correct = sum(bool(r["correct"]) for r in g.values())
        vids = len({r["_v"] for r in g.values()})
        sel_for_ref = s
        if d in ("ALL_TEST", "LECTURE_TEST"):
            def merged(fn):
                m = {}
                for t in (tests if d == "ALL_TEST" else tests[1:]):
                    x = fn(t, sel_for_ref.replace("ALL", t).replace("LECTURE", t), w) if fn is ref_best else fn(t, sel_for_ref)
                    if x is None:
                        return None
                    m.update(x)
                return m
            rs, rb = merged(ref_start), merged(ref_best)
        else:
            rs, rb = ref_start(d, s), ref_best(d, s, w)
        dd = "cgbench_test" if d in ("ALL_TEST", "LECTURE_TEST") else d
        ss = s.replace("ALL", "cgbench_test").replace("LECTURE", "cgbench_test") if d in ("ALL_TEST", "LECTURE_TEST") else s
        self_start = (dd, p, ss, w) == ref_start_key(dd, ss)
        self_best = (dd, p, ss, w) == ref_best_key(dd, ss, w)
        cmp_s = paired(g, rs) if rs is not None and not self_start else None
        cmp_b = paired(g, rb) if rb is not None and not self_best else None
        res = round(100 * correct / n, 1)
        meaning = f"{desc['name']} answered {res}% of {n} {bench} {qset} questions correctly ({desc['model'].split('-')[-1]} answer model)."
        if cmp_s:
            meaning += f" {cmp_s[0]:+.1f} points vs the starting framework ({'confirmed' if cmp_s[2] != 'no' else 'not confirmed'})."
        if cmp_b:
            meaning += f" {cmp_b[0]:+.1f} points vs the best framework ({'confirmed' if cmp_b[2] != 'no' else 'not confirmed'})."
        out.append(dict(stage=stage_of(d, p, s, w) if d not in ("ALL_TEST", "LECTURE_TEST") else (11 if stage_of("cgbench_test", p, s, w) in (7, 8, 9, 10) else stage_of("cgbench_test", p, s, w)),
                        exp="multiple-choice accuracy", desc=desc, bench=bench, vtype=vtype, qset=qset, n=n, vids=vids,
                        measure="multiple-choice accuracy (answer letter correct)", res=res, ct=f"{correct} / {n}",
                        cs=cmp_s, cb=cmp_b, meaning=meaning, sort=(d, p, s, w)))
        ev = [r["chosen_hits_evidence"] for r in g.values() if r.get("chosen_hits_evidence") is not None]
        if ev and d.startswith("cgbench"):
            k = int(sum(ev))
            out.append(dict(stage=out[-1]["stage"], exp="evidence hit", desc=desc, bench=bench, vtype=vtype, qset=qset,
                            n=len(ev), vids=vids, measure="evidence hit: the shown pictures include the human-marked answer moment",
                            res=round(100 * k / len(ev), 1), ct=f"{k} / {len(ev)}", cs=None, cb=None,
                            meaning=f"For {round(100 * k / len(ev), 1)}% of these questions, {desc['name']} showed a picture from the seconds humans marked as the answer.",
                            sort=(d, p, s, w, "ev")))
    return out


def open_rows():
    path = RUN / "open_judged.jsonl"
    if not path.exists():
        return []
    vid = {}
    for line in Path("data/cgbench/qa.jsonl").read_text(encoding="utf-8").splitlines():
        r = json.loads(line)
        vid[r["qa_id"]] = r["video_id"]
    g = collections.defaultdict(dict)
    for line in path.read_text(encoding="utf-8").splitlines():
        r = json.loads(line)
        if r["open_correct"] is not None:
            g[(r["dataset"], r["pool"], r["selector"], r["words"])][(r["dataset"], r["qa_id"])] = {
                "correct": bool(r["open_correct"]), "_v": vid.get(r["qa_id"], r["qa_id"])}
    out = []
    for (d, p, s, w), x in g.items():
        if len(x) < 20:
            continue
        desc = describe(p, s, w)
        bench, vtype, qset = DATA[d]
        ref = g.get((d, "v2", "clip", 120))
        cs = paired(x, ref) if ref is not None and ref is not x else None
        k = sum(v["correct"] for v in x.values())
        res = round(100 * k / len(x), 1)
        out.append(dict(stage=7 if qset == "test" else (4 if w in (0, 120) and "tp=" not in s else 8),
                        exp="open-ended correctness", desc=desc, bench=bench, vtype=vtype, qset=qset, n=len(x),
                        vids=len({v["_v"] for v in x.values()}),
                        measure="open-ended: the model writes its own answer (no options); AI judge marks it correct",
                        res=res, ct=f"{k} / {len(x)}", cs=cs, cb=None,
                        meaning=f"{desc['name']} wrote a correct free answer for {res}% of {len(x)} {bench} {qset} questions.",
                        sort=(d, p, s, w, "open")))
    return out


def manual_rows():
    """Results measured by other scripts (transcribed from the committed reports)."""
    R = []

    def add(stage, exp, name, path_a, path_b, picker, pics, text, model, trained, bench, vtype, qset, n, vids, measure,
            res, ct="", cs=None, cb=None, meaning=""):
        R.append(dict(stage=stage, exp=exp, desc=dict(name=name, path_a=path_a, path_b=path_b, picker=picker, pics=pics,
                                                      text=text, model=model, trained=trained),
                      bench=bench, vtype=vtype, qset=qset, n=n, vids=vids, measure=measure, res=res, ct=ct, cs=cs, cb=cb,
                      meaning=meaning, sort=("m", len(R))))
    lvb = ("LongVideoBench", "everyday videos (8 s - 60 min)", "validation (also used to build v1)", 440, 252)
    for name, picker, res, note in [
            ("v1: transcript only", "no pictures", 32.5, "Answered from the transcript alone."),
            ("v1: evenly spaced pictures", "pictures spread evenly over the video", 39.3, ""),
            ("v1: transcript-retrieved moments", "start / middle / end of the best BM25 windows", 36.4, ""),
            ("v1: MobileCLIP top pictures", "scout scores all candidates, take the most question-like", 41.6,
             "Best v1 strategy; +9.1 over transcript only (confirmed)."),
            ("v1: cost-aware controller", "controller: look / expand / stop by gain minus time cost", 37.5,
             "The v1 research design; 4.1 below MobileCLIP (confirmed).")]:
        add(1, "v1 benchmark", name, "BM25 windows (v1)", "v1 candidate moments", picker, "up to 8", "BM25 excerpt, 120 words",
            "Qwen3-VL-2B", "none (controller rule hand-set)", *lvb, "multiple-choice accuracy", res,
            meaning=f"{name} answered {res}% of 440 LongVideoBench questions correctly. {note}")
    for name, res in [("v1: transcript only", 44.1), ("v1: evenly spaced pictures", 33.4), ("v1: transcript-retrieved moments", 32.0),
                      ("v1: MobileCLIP top pictures", 32.3), ("v1: cost-aware controller", 29.8)]:
        add(1, "v1 open-ended", name, "BM25 windows (v1)", "v1 candidate moments", "as above", "up to 8", "BM25 excerpt, 120 words",
            "Qwen3-VL-2B", "none", *lvb, "open-ended correct (local Phi-4-mini judge, never checked against humans)", res,
            meaning=f"{res}% judged correct; the judge was not validated, so these numbers were not used for claims.")
    add(1, "v1 research idea", "v1 controller trained to look when important speech is damaged", "BM25 windows (v1)",
        "v1 candidate moments", "trained controller (tiny network; Qwen3-0.6B + LoRA)", "up to 8", "damaged transcript",
        "Qwen3-VL-2B", "v1 controller", "LongVideoBench", "everyday videos", "pilot train", 26, 12,
        "share of cases where damaging the important speech changed the answer", 13.0,
        meaning="Only 13% (train) / 4% (dev, 14 questions) of damage cases mattered; the controller could not learn when to look. Idea not supported.")
    for bench, k, n in [("LongVideoBench", 2, 50), ("Video-MME", 2, 50), ("EduVidQA", 0, 43), ("TVQA", 2, 50)]:
        add(1, "v1 research check", "Questions where losing one speech piece breaks the answer", "-", "-", "-", "-",
            "clean vs damaged transcript", "Qwen3-VL-2B", "none", bench, "", "sample", n, "",
            "share of questions where the answer depends on one local piece of speech", round(100 * k / n, 1), f"{k} / {n}",
            meaning=f"Only {k} of {n} {bench} questions fit the v1 research idea, so it could not be tested.")
    mit = ("MIT lecture questions (self-made)", "lecture videos (MIT OpenCourseWare)", "each lecture tested by models that never saw it", 197, 20)
    for name, picker, trained, res, note in [
            ("v2: transcript only", "no pictures", "none", 15.2, "Below chance because of an answer-position bias in the generated questions."),
            ("v2: MobileCLIP top-4", "top pictures by match to the question", "none", 34.5, ""),
            ("v2: MobileCLIP + diversity", "MMR rule on the question", "none", 36.0, ""),
            ("v2: learned Head B with screen text", "v2 Head B (trained)", "v2 Head B", 34.7, "No gain over simple rules."),
            ("v2: learned Head B without screen text", "v2 Head B (trained)", "v2 Head B", 32.8, ""),
            ("v2: learned Head A frame credit", "v2 Head A (trained)", "v2 Head A", 30.5, "Confirmed worse than the zero-shot rule."),
            ("v2: zero-shot transcript relevance", "top pictures by speech score (untrained)", "none", 38.6, "Best v2 picker."),
            ("v2: learned Head A text credit", "v2 Head A text credit (trained)", "v2 Head A", 37.1, "")]:
        add(2, "v2 main experiment", name, "speech rule inside BM25 windows", "scan inside BM25 windows", picker, "4",
            "BM25 excerpt, 120 words", "Qwen3-VL-2B", trained, *mit, "multiple-choice accuracy", res,
            meaning=f"{name}: {res}% of 197 MIT lecture questions. {note}")
    add(2, "v2 pilot", "v2 learned pickers (pilot)", "speech rule", "scan inside windows", "learned pickers", "4", "BM25 120 words",
        "Qwen3-VL-2B", "v2 heads", "MIT lecture questions (self-made)", "lectures", "pilot", 132, 5, "accuracy vs MobileCLIP top-4",
        "", meaning="Every learned picker was within noise of MobileCLIP top-4 (not supported).")
    add(2, "v2 follow-up", "Learned 4th-picture completion", "speech rule", "scan inside windows", "top-3 + learned 4th picture", "4",
        "BM25 120 words", "Qwen3-VL-2B", "completion scorer", *mit, "accuracy difference vs the rule's own 4th picture", -1.0,
        cs=None, meaning="-1.0 points (interval -2.5 to 0.0): refuted.")
    add(2, "v2 follow-up", "Picture scan spread evenly over windows", "speech rule", "balanced scan inside windows",
        "top pictures by match to the question", "4", "BM25 120 words", "Qwen3-VL-2B", "none", *mit, "multiple-choice accuracy", 39.1,
        meaning="39.1% vs 34.5% with the original scan (+4.6, interval -0.5 to +10.4): not confirmed.")
    add(2, "v2 follow-up", "Option-rotation scoring, MobileCLIP top-4", "speech rule", "scan inside windows", "top pictures by match",
        "4", "BM25 120 words", "Qwen3-VL-2B", "none", *mit, "accuracy averaged over the 4 option orders", 49.7,
        meaning="49.7% averaged over option rotations vs 34.3% for transcript only: fixes the answer-position bias.")
    for name, res in [("v2 scan on LongVideoBench: original", 43.0), ("v2 scan on LongVideoBench: spread evenly", 42.5)]:
        add(2, "v2 transfer check", name, "speech rule", "scan inside windows", "MobileCLIP top-4", "4", "BM25 120 words",
            "Qwen3-VL-2B", "none", "LongVideoBench", "everyday videos", "validation", 440, 252, "multiple-choice accuracy", res,
            meaning="The MIT finding did not transfer to a public benchmark; v3 moved to public benchmarks.")
    for name, res in [("No pictures (transcript only)", 0.587), ("4 evenly spaced pictures", 0.493), ("4 MobileCLIP-chosen pictures", 0.498)]:
        add(3, "v3 benchmark check", name, "-", "-", name, "0" if "No" in name else "4", "whole transcript at the question time",
            "Qwen3-VL-2B", "none", "EduVidQA", "lecture videos, open-ended", "official train (checked only)", 74, 20,
            "FactQA precision (share of supported facts, 0-1, not a percentage)", res,
            meaning="Pictures made the small model's lecture answers worse; EduVidQA was not used for picture claims.")
    for name, res, pics in [("Question only", 28.9, "0"), ("Transcript only", 40.2, "0"), ("4 evenly spaced pictures + transcript", 52.6, "4"),
                            ("4 MobileCLIP pictures + transcript", 56.7, "4")]:
        add(3, "v3 benchmark check", name, "-", "-", name, pics, "whole transcript (2,000 words max)", "Qwen3-VL-2B", "none",
            "Video-MMMU Perception", "lectures, screen-reading questions", "dev", 97, 97, "multiple-choice accuracy", res,
            meaning=f"{res}%: pictures help on lectures (+16.5 for MobileCLIP pictures over transcript only, confirmed).")
    for name, res, pics in [("Question only", 15.9, "0"), ("Transcript only (BM25 300 words)", 15.9, "0"), ("4 evenly spaced pictures", 17.1, "4"),
                            ("4 MobileCLIP pictures (dense, whole video)", 40.2, "4"), ("4 pictures from the human-marked answer seconds", 63.4, "4")]:
        add(3, "v3 benchmark check", name, "-", "-", name, pics, "BM25 300 words" if pics == "4" else name, "Qwen3-VL-2B", "none",
            "CG-Bench", "long everyday videos (10-80 min)", "dev", 82, 8, "multiple-choice accuracy", res,
            meaning=f"{res}%." + (" Ceiling for the 2B model with perfect pictures." if "human" in name else ""))
    ha = ("CG-Bench", "long everyday videos (10-80 min)", "dev", 164, 17)
    for name, res, trained, note in [
            ("Speech rule inside BM25 windows (best framework's Path A)", 25.0, "none", ""),
            ("MiniLM over the whole transcript (untrained)", 28.0, "none", "Reference for Head A."),
            ("Fine-tuned MiniLM", 8.0, "fine-tuned MiniLM", "About 8%: training destroyed the ranking."),
            ("Larger rerankers, untrained (best: BGE-base / BGE-v2-m3)", 31.1, "none", "Model size barely helps."),
            ("Light trained network, with EduVidQA weak data", 32.7, "Head A draft", ""),
            ("Light trained network, human data only", 32.1, "Head A draft", ""),
            ("Light network + lines vote", 35.2, "Head A draft", ""),
            ("Light network + learned offsets (original idea)", 31.7, "Head A draft", "Offsets did not help."),
            ("Head A (chosen): lines vote + BGE facts, 3 copies averaged", 36.6, "Head A", "+8.5 over 28.0 (interval +4.1 to +13.1), confirmed.")]:
        add(5, "Head A method screen", name, name, "-", "-", "6 moments", "-", "-", trained, *ha,
            "evidence hit: at least one of 6 moments inside the human-marked answer seconds", res,
            meaning=f"{res}% of 164 new CG-Bench questions. Trained on 877 CG-Bench questions (83 videos). {note}")
    hb = ("CG-Bench + Video-MMMU (both tracks)", "long videos and lectures", "dev", 278, 113)
    for name, diff, note in [("Option-evidence set rule (untrained)", -1.4, ""), ("Head B trained on human-marked seconds", -1.4, ""),
                             ("Head B + option term (chosen Head B)", -0.4, ""), ("Head B trained on answer-model feedback", -3.6, "Worse on lectures."),
                             ("Head B feedback + option term", -4.0, ""), ("Confidence choice over 4 picture sets", 1.1, "Interval -1.7 to +3.8.")]:
        add(6, "Head B method screen (pooled dev)", name, "speech rule", "hybrid", name, "4", "none", "Qwen3-VL-2B",
            "Head B" if "Head B" in name else "none", *hb, "accuracy difference to the MMR rule (points, pooled over 278 dev questions)",
            diff, cb=(diff, "", "no"),
            meaning=f"{diff:+.1f} points vs the MMR rule. Head B trained on 373 CG-Bench questions (37 videos). {note}")
    cov = ("CG-Bench", "long everyday videos (10-80 min)", "dev", 82, 8)
    for name, res in [("v2 BM25 excerpt, 120 words", 28.0), ("No words", 0.0), ("Meaning-based search, 120 words", 50.0),
                      ("Meaning-based search, 300 words", 62.2), ("Speech around the pictures, 300 words", 48.8),
                      ("Speech around pictures + search, 300 words", 65.9), ("Speech around pictures + search, 600 words (chosen)", 68.3),
                      ("Whole transcript, max 1,500 words", 65.9), ("Whole transcript, max 4,000 words", 79.3), ("Whole transcript, no cap", 80.5)]:
        add(8, "speech kept", name, "speech rule", "hybrid", "MMR rule on question + options", "4", name, "-", "none", *cov,
            "speech kept: the given words include speech within 10 s of the human-marked answer moment", res,
            meaning=f"{res}% of questions: the text given to the answer model contains the speech near the answer.")
    return R


def main() -> None:
    groups = load_rows()
    rows = manual_rows() + v3_rows(groups) + open_rows()
    order = {"ALL_TEST": 9, "LECTURE_TEST": 10}
    rows.sort(key=lambda r: (r["stage"], r["sort"][0] == "m" and -1 or 0,
                             str(order.get(r["sort"][0], r["sort"][0])), str(r["sort"][1:])))
    table = []
    for i, r in enumerate(rows, 1):
        d, cs, cb = r["desc"], r["cs"], r["cb"]
        table.append([i, STAGES[r["stage"]], r["exp"], d["name"], d["path_a"], d["path_b"], d["picker"], d["pics"], d["text"],
                      d["model"], d["trained"], r["bench"], r["vtype"], r["qset"], r["n"], r["vids"], r["measure"], r["res"], r["ct"],
                      cs[0] if cs else "", cs[1] if cs else "", cs[2] if cs else "",
                      cb[0] if cb else "", cb[1] if cb else "", cb[2] if cb else "", r["meaning"]])
    OUT.mkdir(parents=True, exist_ok=True)
    with open(OUT / "all_results.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(COLS)
        w.writerows(table)
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    ws.title = "All results"
    ws.append(COLS)
    for row in table:
        ws.append(row)
    widths = [6, 22, 22, 34, 30, 30, 30, 9, 34, 13, 16, 26, 26, 14, 10, 8, 40, 10, 12, 14, 18, 14, 14, 18, 14, 70]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    head = PatternFill("solid", fgColor="DDE6EE")
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True), head, Alignment(wrap_text=True, vertical="top")
    green, red = PatternFill("solid", fgColor="E3F4E6"), PatternFill("solid", fgColor="F8E1DF")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for idx in (21, 24):
            v = row[idx].value
            if v == "yes (gain)":
                row[idx].fill = green
            elif v == "yes (loss)":
                row[idx].fill = red
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    gl = wb.create_sheet("How to read")
    for line in [
            ["Column", "Meaning"],
            ["Stage", "the experiments in the order they were done (1 = first)"],
            ["System (plain name)", "Best framework = untrained tuned rules (speech rule in Path A, hybrid picture search, MMR picker). Head A = trained speech path. Head B = trained picker. Speech text = up to 600 transcript words given to the answer model."],
            ["Question set", "train = used to train the heads; dev = used to choose settings; test = locked, used once at the end"],
            ["Result (%)", "the value of 'What is measured' (FactQA rows use a 0-1 score, Head B screen rows are point differences)"],
            ["vs ... difference", "result minus the reference on the same questions, in percentage points"],
            ["95% interval / confirmed?", "paired bootstrap over videos; 'yes' when the interval excludes zero"],
            ["Answer model", "frozen vision-language model that reads the pictures and answers (2B small, 4B larger)"],
            ["Pooled rows (stage 11)", "the same system over all 459 test questions, or over the 404 lecture test questions"]]:
        gl.append(line)
    gl.column_dimensions["A"].width, gl.column_dimensions["B"].width = 26, 120
    wb.save(OUT / "all_results.xlsx")
    css = ("@page{size:A3 landscape;margin:8mm} body{font-family:'Segoe UI',Arial;font-size:6.4pt;color:#1c2330}"
           "h1{font-size:13pt;margin:0 0 4px} p{margin:0 0 6px;font-size:8pt}"
           "table{border-collapse:collapse;width:100%} th,td{border:1px solid #d4dae2;padding:1.5px 3px;vertical-align:top}"
           "th{background:#dde6ee;position:sticky;top:0} tr{break-inside:avoid} .g{background:#e3f4e6} .r{background:#f8e1df}")
    hdr = "".join(f"<th>{c}</th>" for c in COLS)
    body = []
    for row in table:
        cells = []
        for j, v in enumerate(row):
            cls = ' class="g"' if v == "yes (gain)" else (' class="r"' if v == "yes (loss)" else "")
            cells.append(f"<td{cls}>{v}</td>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    html = (f"<!doctype html><html><head><meta charset=utf-8><title>All Results</title><style>{css}</style></head><body>"
            f"<h1>All results, first experiment to last</h1><p>One result per row. Best framework = untrained tuned rules; "
            f"Head A = trained speech path; Head B = trained picker; speech text = up to 600 transcript words for the answer model. "
            f"dev = used to choose settings, test = locked and used once. Confirmed = the 95% interval excludes zero (green gain, red loss).</p>"
            f"<table><tr>{hdr}</tr>{''.join(body)}</table></body></html>")
    (OUT / "all_results.html").write_text(html, encoding="utf-8")
    pdf = (OUT / "all_results.pdf").resolve()
    exe = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
    if Path(exe).exists():
        subprocess.run([exe, "--headless", "--disable-gpu", "--no-pdf-header-footer", f"--print-to-pdf={pdf}",
                        (OUT / "all_results.html").resolve().as_uri()], check=True, timeout=180)
    print(f"{len(table)} rows -> {OUT / 'all_results.xlsx'}, .csv, .pdf")


if __name__ == "__main__":
    main()
